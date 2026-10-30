"""
Fixtures compartidas por todas las pruebas 🧪
Cada prueba usa su PROPIA base SQLite temporal (tmp_path), cargada con el MISMO pipeline
ETL que los datos reales, pero desde un Excel PEQUEÑO falso que se arma aquí con openpyxl
a partir de data/precios_ejemplo.csv (precios ficticios). Así no hace falta internet, ni el
Excel de 64 MB del DANE, ni el .env real (Telegram queda "sin configurar").
"""
import csv

import openpyxl
import pytest

from agente import config
from scripts.actualizar_sipsa import actualizar

ENCABEZADO = ("Año ", "Mes", "Código departamento", "Nombre departamento ", "Código municipio",
              "Nombre municipio ", "Nombre del producto", "Presentación del producto", "Precio promedio ")


def filas_ejemplo() -> list[tuple]:
    """Filas del CSV de ejemplo en el orden de columnas del Excel del DANE."""
    with open(config.RUTA_PRECIOS, encoding="utf-8") as f:
        return [(int(r["anio"]), r["mes"], r["cod_departamento"], r["departamento"], r["cod_municipio"],
                 r["municipio"], r["producto"], r["presentacion"], float(r["precio"]))
                for r in csv.DictReader(f)]


def crear_excel_falso(ruta, filas: list[tuple]) -> None:
    """Excel con la misma forma que el del DANE: títulos en las filas 1-8, encabezado en la
    9, datos desde la 10, meses en texto, espacios sobrantes, filas vacías y notas al pie.
    Trae además una hoja '1.4' que el ETL NO debe leer (solo lee HOJAS_SIPSA)."""
    libro = openpyxl.Workbook()
    libro.active.title = "Índice"
    hoja = libro.create_sheet("1.3")
    hoja["J1"] = "Regresar al índice"
    hoja["A3"] = "SIPSA-I Serie histórica de precio promedio de mercado por municipio"
    hoja["A5"] = "1.3. Fertilizantes, enmiendas y acondicionadores de suelo"
    for col, titulo in enumerate(ENCABEZADO, start=1):
        hoja.cell(row=9, column=col, value=titulo)
    fila_excel = 10
    for i, fila in enumerate(filas):
        anio, mes, cod_dep, dep, cod_mun, mun, producto, presentacion, precio = fila
        if i % 2:  # el DANE deja espacios sobrantes en algunos textos
            mes, dep, mun = mes + "  ", dep + " ", mun + " "
        for col, valor in enumerate((anio, mes, cod_dep, dep, cod_mun, mun, producto, presentacion, precio), 1):
            hoja.cell(row=fila_excel, column=col, value=valor)
        fila_excel += 2 if i == 5 else 1  # una fila vacía en medio de los datos
    fila_excel += 2  # filas vacías al final y luego las notas, en la columna del año
    for nota in ("Fuente: DANE, SIPSA-I", "Actualizado el 20 de octubre de 2026"):
        hoja.cell(row=fila_excel, column=1, value=nota)
        fila_excel += 1

    otra = libro.create_sheet("1.4")
    for col, titulo in enumerate(ENCABEZADO, start=1):
        otra.cell(row=9, column=col, value=titulo)
    for col, valor in enumerate((2026, "Septiembre", "68", "Santander", "68001", "Bucaramanga",
                                 "Urea 46%", "50 kilogramos", 999999), 1):
        otra.cell(row=10, column=col, value=valor)
    libro.save(ruta)


@pytest.fixture
def excel_falso(tmp_path):
    ruta = tmp_path / "sipsa_falso.xlsx"
    crear_excel_falso(ruta, filas_ejemplo())
    return ruta


@pytest.fixture(autouse=True)
def bd_temporal(tmp_path, monkeypatch, excel_falso):
    ruta = tmp_path / "insumos_test.db"
    actualizar(excel_falso, ruta, hojas=["1.3"], anios_historia=2)
    monkeypatch.setattr(config, "RUTA_BD", ruta)
    monkeypatch.setattr(config, "UMBRAL_PCT_MENSUAL", 2.0)
    monkeypatch.setattr(config, "MESES_DATO_VIGENTE", 3)
    monkeypatch.setattr(config, "TELEGRAM_BOT_TOKEN", "")
    monkeypatch.setattr(config, "TELEGRAM_CHAT_ID", "")
    monkeypatch.delenv("SIMULAR_FALLA_TRM", raising=False)
    return ruta
