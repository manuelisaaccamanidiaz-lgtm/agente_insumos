"""
actualizar_sipsa.py — Pipeline ETL de precios reales del DANE 🏭
================================================================
  EXTRAER     ─► descarga el Excel de SIPSA-Insumos (si no está o está viejo)
  TRANSFORMAR ─► lee las hojas configuradas, deja solo los insumos del catálogo y los
                 últimos años, convierte "Enero " → 1 y limpia espacios
  CARGAR      ─► guarda en data/insumos.db (tabla precios_sipsa), en una transacción

Es IDEMPOTENTE: correrlo dos veces no duplica filas (índice único). Si el DANE corrige
un precio, la fila se actualiza.

Uso:
    python scripts/actualizar_sipsa.py            (descarga solo si hace falta)
    python scripts/actualizar_sipsa.py --forzar   (descarga siempre)
    python scripts/actualizar_sipsa.py --dias 7   (re-descarga si el archivo tiene > 7 días)

Formato del Excel (verificado con el de 2021-2026): filas 1-8 son títulos, la fila 9 es
el encabezado y los datos empiezan en la 10. Al final vienen notas ("Fuente: DANE…")
en la columna del año: se ignoran porque no son un año.
"""
import argparse
import sys
import time
from contextlib import closing
from pathlib import Path

import openpyxl
import requests

# Permite correr el script directamente (python scripts/actualizar_sipsa.py) e importar `agente`
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from agente import config  # noqa: E402
from agente.catalogo import PRESENTACIONES, normalizar  # noqa: E402
from agente.datos import conectar, crear_tablas, insertar_precios_sipsa  # noqa: E402

FILA_ENCABEZADO = 9
FUENTE = "DANE-SIPSA"
MESES = {m: i for i, m in enumerate(
    ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
     "septiembre", "octubre", "noviembre", "diciembre"], start=1)}
MESES["setiembre"] = 9  # variante que a veces aparece en documentos oficiales
TAMANO_PEDAZO = 1024 * 1024  # la descarga se escribe a disco de a 1 MB


class DescargaFallida(RuntimeError):
    """No se pudo descargar el Excel y tampoco hay una copia local para usar."""


# ---------- 1. EXTRAER ----------
def necesita_descarga(ruta: Path, dias: int, forzar: bool = False) -> bool:
    """True si no hay copia local, si tiene más de `dias` días o si se pidió --forzar."""
    if forzar or not ruta.exists():
        return True
    return (time.time() - ruta.stat().st_mtime) / 86400 > dias


def descargar(url: str, destino: Path) -> Path:
    """Descarga en streaming a un archivo temporal (.part) y solo al final lo renombra:
    si la descarga se corta, la copia local anterior queda intacta."""
    temporal = destino.with_name(destino.name + ".part")
    destino.parent.mkdir(parents=True, exist_ok=True)
    try:
        with requests.get(url, stream=True, timeout=config.TIMEOUT_DESCARGA) as resp:
            resp.raise_for_status()
            total = int(resp.headers.get("content-length") or 0)
            bajado = 0
            with open(temporal, "wb") as f:
                for pedazo in resp.iter_content(chunk_size=TAMANO_PEDAZO):
                    f.write(pedazo)
                    bajado += len(pedazo)
                    avance = f" ({bajado / total:.0%})" if total else ""
                    print(f"\r   ⬇️  {bajado / 1e6:,.1f} MB{avance}", end="", flush=True)
        print()
        temporal.replace(destino)
        return destino
    finally:
        temporal.unlink(missing_ok=True)


def extraer(url: str, ruta: Path, dias: int, forzar: bool = False) -> Path:
    """Devuelve la ruta del Excel a usar. Si la descarga falla pero hay copia local,
    usa la local y avisa. Si no hay nada, lanza DescargaFallida."""
    if not necesita_descarga(ruta, dias, forzar):
        edad = (time.time() - ruta.stat().st_mtime) / 86400
        print(f"📄 Uso la copia local ({edad:.0f} días, límite {dias}): {ruta.name}")
        return ruta
    print(f"🌐 Descargando {url}")
    try:
        return descargar(url, ruta)
    except requests.RequestException as e:
        detalle = getattr(getattr(e, "response", None), "status_code", None) or type(e).__name__
        if ruta.exists():
            print(f"⚠️  La descarga falló ({detalle}). Sigo con la copia local: {ruta.name}")
            return ruta
        raise DescargaFallida(
            f"No se pudo descargar el Excel del DANE ({detalle}) y no hay copia local en {ruta}. "
            "Revisa el internet o SIPSA_URL en el .env (el año del nombre cambia cada año)."
        ) from e


# ---------- 2. TRANSFORMAR ----------
def mes_a_numero(texto) -> int | None:
    """'Enero ' -> 1, 'SEPTIEMBRE' -> 9. None si no es un mes."""
    return MESES.get(normalizar(texto)) if isinstance(texto, str) else None


def _texto(valor) -> str:
    """Celda → texto sin espacios sobrantes ('Santander ' -> 'Santander', None -> '')."""
    return " ".join(str(valor).split()) if valor is not None else ""


def _filas_de_hoja(hoja) -> list[tuple]:
    """Filas de datos de una hoja (desde la fila 10). Valida el encabezado de la fila 9
    para avisar claro si el DANE cambió el formato del archivo."""
    filas = hoja.iter_rows(min_row=FILA_ENCABEZADO, values_only=True)
    encabezado = next(filas, None) or ()
    if not normalizar(_texto(encabezado[0] if encabezado else "")).startswith("ano"):
        raise ValueError(f"La hoja '{hoja.title}' no tiene el encabezado 'Año' en la fila "
                         f"{FILA_ENCABEZADO}: el DANE pudo cambiar el formato del archivo.")
    return filas


def transformar(ruta_excel: Path, hojas: list[str], anios_historia: int) -> tuple[list[dict], int]:
    """Lee el Excel y devuelve (registros limpios, filas de datos leídas).
    Se quedan solo: insumos del catálogo con su presentación, y los `anios_historia`
    años más recientes del archivo (ej. 2 → 2026 y 2025)."""
    libro = openpyxl.load_workbook(ruta_excel, read_only=True, data_only=True)
    registros, leidas = [], 0
    try:
        for nombre in hojas:
            if nombre not in libro.sheetnames:
                print(f"⚠️  La hoja '{nombre}' no existe en el Excel; la salto.")
                continue
            for fila in _filas_de_hoja(libro[nombre]):
                anio, mes, cod_dep, dep, cod_mun, mun, producto, presentacion, precio = (
                    tuple(fila) + (None,) * 9)[:9]
                if not isinstance(anio, (int, float)):  # filas vacías y notas al pie
                    continue
                leidas += 1
                producto, presentacion = _texto(producto), _texto(presentacion)
                numero_mes = mes_a_numero(mes)
                if (PRESENTACIONES.get(producto) != presentacion or numero_mes is None
                        or not isinstance(precio, (int, float))):
                    continue
                registros.append({
                    "anio": int(anio), "mes": numero_mes, "periodo": f"{int(anio)}-{numero_mes:02d}",
                    "cod_departamento": _texto(cod_dep), "departamento": _texto(dep),
                    "cod_municipio": _texto(cod_mun), "municipio": _texto(mun),
                    "producto": producto, "presentacion": presentacion,
                    "precio_cop": float(precio), "fuente": FUENTE,
                })
    finally:
        libro.close()  # en modo read_only el archivo queda abierto hasta cerrarlo
    if registros:
        anio_minimo = max(r["anio"] for r in registros) - anios_historia + 1
        registros = [r for r in registros if r["anio"] >= anio_minimo]
    return registros, leidas


# ---------- 3. CARGAR ----------
def cargar(registros: list[dict], ruta_bd: Path) -> tuple[int, int]:
    """Guarda los registros en precios_sipsa. Devuelve (nuevas, actualizadas)."""
    Path(ruta_bd).parent.mkdir(parents=True, exist_ok=True)
    with closing(conectar(ruta_bd, crear=True)) as conexion:
        crear_tablas(conexion)
        return insertar_precios_sipsa(conexion, registros)


def actualizar(ruta_excel: Path, ruta_bd: Path, hojas: list[str], anios_historia: int) -> dict:
    """Transformar + cargar. Devuelve el resumen (también lo usan las pruebas)."""
    registros, leidas = transformar(ruta_excel, hojas, anios_historia)
    nuevas, actualizadas = cargar(registros, ruta_bd)
    periodos = sorted({r["periodo"] for r in registros})
    return {
        "filas_leidas": leidas,
        "registros_validos": len(registros),
        "insertadas": nuevas,
        "actualizadas": actualizadas,
        "desde": periodos[0] if periodos else None,
        "hasta": periodos[-1] if periodos else None,
        "productos": {p: sum(r["producto"] == p for r in registros)
                      for p in sorted({r["producto"] for r in registros})},
    }


def main(argv: list[str] | None = None) -> int:
    sys.stdout.reconfigure(encoding="utf-8")  # emojis sin errores en la consola de Windows
    parser = argparse.ArgumentParser(description="Actualiza los precios reales del DANE (SIPSA-Insumos).")
    parser.add_argument("--forzar", action="store_true", help="descarga el Excel aunque haya copia local")
    parser.add_argument("--dias", type=int, default=config.DIAS_ACTUALIZAR_SIPSA,
                        help="re-descarga si la copia local tiene más de N días (por defecto %(default)s)")
    args = parser.parse_args(argv)

    try:
        ruta_excel = extraer(config.SIPSA_URL, config.RUTA_SIPSA, args.dias, args.forzar)
    except DescargaFallida as e:
        print(f"❌ {e}")
        return 1

    inicio = time.perf_counter()
    print(f"🔄 Leyendo hojas {', '.join(config.HOJAS_SIPSA)} (últimos {config.ANIOS_HISTORIA} años)...")
    resumen = actualizar(ruta_excel, config.RUTA_BD, config.HOJAS_SIPSA, config.ANIOS_HISTORIA)

    print(f"🗄️  BD: {config.RUTA_BD} (tabla precios_sipsa) · {time.perf_counter() - inicio:.1f} s")
    print(f"✅ {resumen['filas_leidas']:,} filas leídas · {resumen['registros_validos']:,} del catálogo "
          f"· {resumen['insertadas']:,} nuevas · {resumen['actualizadas']:,} actualizadas")
    print(f"📅 Periodos: {resumen['desde']} → {resumen['hasta']}")
    print("🌱 Productos cargados:")
    for producto, n in resumen["productos"].items():
        print(f"   - {producto}: {n:,} registros")
    return 0


if __name__ == "__main__":
    sys.exit(main())
